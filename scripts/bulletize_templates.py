#!/usr/bin/env python3
"""Tambah bullet poin di bawah 'ADALAH BENAR' pada template surat .docx."""
import re
import sys
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

STD = ("Bahwa saudara/saudari tersebut di atas adalah penduduk "
       "Desa Ngadri Kecamatan Binangun Kabupaten Blitar.")
BULLET = "•\t"

PLACEHOLDER_ONLY = re.compile(r"(\$\{\w+\}[\s.,;]*)+$")


def is_placeholder_only(text):
    return bool(PLACEHOLDER_ONLY.fullmatch(text.strip()))


def is_sub_heading(text):
    letters = re.sub(r"[^A-Za-z]", "", text)
    return bool(letters) and letters.isupper()


def set_hanging_indent(p):
    pf = p.paragraph_format
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = -Cm(0.75)


def prefix_bullet(p):
    r = p.runs[0]
    r.text = BULLET + r.text
    set_hanging_indent(p)


def replace_text(p, new_text):
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = new_text
    set_hanging_indent(p)


def insert_std_before(ref, fmt_src=None):
    fmt_src = fmt_src if fmt_src is not None else ref
    new_p = ref.insert_paragraph_before()
    ref_pPr = fmt_src._p.find(qn("w:pPr"))
    if ref_pPr is not None:
        new_p._p.insert(0, deepcopy(ref_pPr))
    run = new_p.add_run(BULLET + STD)
    if fmt_src.runs:
        ref_rPr = fmt_src.runs[0]._r.find(qn("w:rPr"))
        if ref_rPr is not None:
            run._r.insert(0, deepcopy(ref_rPr))
    set_hanging_indent(new_p)


def bulletize(path):
    d = Document(path)
    paras = d.paragraphs

    heading_idx = None
    for i, p in enumerate(paras):
        if p.text.strip().upper().startswith("ADALAH BENAR"):
            heading_idx = i
            break
    if heading_idx is None:
        return "tanpa-ADALAH-BENAR"

    points = []
    first_any = None
    for p in paras[heading_idx + 1:]:
        t = p.text.strip()
        if t.lower().startswith("demikian"):
            break
        if t and first_any is None:
            first_any = p
        if t and not is_placeholder_only(t) and not is_sub_heading(t):
            points.append(p)

    if not points:
        return "tanpa-poin"

    changed = False
    first = points[0]
    ft = first.text.strip()

    if ft.startswith("•"):
        pass
    elif "penduduk Desa" in ft and "${" not in ft:
        replace_text(first, BULLET + STD)
        changed = True
    else:
        insert_std_before(first_any, fmt_src=first)
        changed = True

    for p in points:
        if p.text.strip().startswith("•"):
            continue
        prefix_bullet(p)
        changed = True

    if changed:
        d.save(path)
    return "OK" if changed else "tidak-berubah"


if __name__ == "__main__":
    for f in sys.argv[1:]:
        print(f"{f}: {bulletize(f)}")
