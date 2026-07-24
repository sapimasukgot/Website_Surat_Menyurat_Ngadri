#!/usr/bin/env python3
"""Putaran 3:
1. Nama TTD pemohon & perangkat selalu rata (vertical-align bottom pada sel ttd).
2. Kalimat pembuka selalu "Kepala Desa Ngadri ..." (statis, tidak ikut jabatan ttd).
3. Jarak antar bullet lebih rapat (6pt) dibanding antar paragraf biasa (12pt).
"""
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Pt


def fix_pembuka(d):
    for p in d.paragraphs:
        if "Yang bertanda tangan" not in p.text:
            continue
        for r in p.runs:
            if "${jabatan_ttd}" in r.text:
                r.text = r.text.replace("${jabatan_ttd}", "Kepala Desa")
                return "pembuka-diganti"
        return "pembuka-sudah-statis"
    return "tanpa-pembuka"


def fix_ttd_valign(d):
    for t in d.tables:
        if "${penandatangan}" not in t._tbl.xml:
            continue
        for row in t.rows:
            for c in row.cells:
                c.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
        return True
    return False


def fix_bullet_spacing(d):
    n = 0
    for p in d.paragraphs:
        if p.text.strip().startswith("•"):
            p.paragraph_format.space_after = Pt(6)
            n += 1
    return n


def process(path):
    d = Document(path)
    a = fix_pembuka(d)
    b = fix_ttd_valign(d)
    c = fix_bullet_spacing(d)
    d.save(path)
    return f"{a} | ttd-bottom:{b} | bullet-6pt:{c}"


if __name__ == "__main__":
    for f in sys.argv[1:]:
        print(f"{f}: {process(f)}")
