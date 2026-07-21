#!/usr/bin/env python3
"""Putaran 2: TTD pemohon di kiri, font Times New Roman, jarak antar paragraf."""
import sys
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "Times New Roman"


def find_kop(d):
    for t in d.tables:
        if "PEMERINTAH" in t._tbl.xml:
            return t
    return None


def find_sig(d, kop):
    for t in d.tables:
        if kop is not None and t._tbl is kop._tbl:
            continue
        if "${penandatangan}" in t._tbl.xml:
            return t
    return None


def set_font_run(r):
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is not None:
        rFonts.set(qn("w:cs"), FONT)
        rFonts.set(qn("w:eastAsia"), FONT)


def apply_font(d, kop):
    for p in d.paragraphs:
        for r in p.runs:
            set_font_run(r)
    for t in d.tables:
        if kop is not None and t._tbl is kop._tbl:
            continue
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        set_font_run(r)
    # style dasar dokumen ikut TNR agar teks baru konsisten
    try:
        st = d.styles["Normal"]
        st.font.name = FONT
    except KeyError:
        pass


def apply_spacing(d):
    """Jarak antar paragraf ±1 baris (12pt) untuk paragraf isi (level body)."""
    for p in d.paragraphs:
        if p.text.strip():
            p.paragraph_format.space_after = Pt(12)


def add_pemohon(d, sig):
    if sig is None:
        return "tanpa-tabel-ttd"
    xml = sig._tbl.xml
    if "Pemohon" in xml or "${nama}" in xml:
        return "sudah-ada"

    row = sig.rows[0]
    if len(row.cells) < 2:
        return "kolom-kurang"

    src = next((c for c in row.cells if "${penandatangan}" in c._tc.xml), None)
    dst = next((c for c in row.cells if c._tc is not src._tc and not c.text.strip()), None)
    if src is None or dst is None:
        return "sel-tidak-cocok"

    texts = ["", "Pemohon,", "", "", "", "${nama}"]
    src_paras = src.paragraphs

    # kosongkan isi sel kiri (biasanya satu paragraf kosong)
    for p in list(dst.paragraphs):
        p._p.getparent().remove(p._p)

    for i, txt in enumerate(texts):
        ref = src_paras[i] if i < len(src_paras) else src_paras[-1]
        new_p = deepcopy(ref._p)
        # buang semua run hasil copy, sisakan pPr
        for r in new_p.findall(qn("w:r")):
            new_p.remove(r)
        dst._tc.append(new_p)
        if txt == "" and not ref.runs:
            continue
        # buat run baru dengan format run referensi
        ref_r = ref.runs[0]._r if ref.runs else None
        from docx.text.paragraph import Paragraph
        para = Paragraph(new_p, dst)
        run = para.add_run(txt)
        if ref_r is not None:
            ref_rPr = ref_r.find(qn("w:rPr"))
            if ref_rPr is not None:
                run._r.insert(0, deepcopy(ref_rPr))
        set_font_run(run)
    return "OK"


def process(path):
    d = Document(path)
    kop = find_kop(d)
    sig = find_sig(d, kop)
    status = add_pemohon(d, sig)
    apply_font(d, kop)
    apply_spacing(d)
    d.save(path)
    return status


if __name__ == "__main__":
    for f in sys.argv[1:]:
        print(f"{f}: {process(f)}")
